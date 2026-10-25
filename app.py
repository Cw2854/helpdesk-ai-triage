"""
 Brightmoor IT digital Desk Portal - Admin web app
--------------------------------
Employees report IT problems through a web form. The AI triages each ticket straight away,
it's saved to a small database, and the admin sees it in the ticket queue.

First time only, create the admin login:
    python set_password.py

Then start the app:
    python app.py
and open http://127.0.0.1:5000 in your browser.
  - Admin sign-in:       http://127.0.0.1:5000/login
  - Employee report form: http://127.0.0.1:5000/report

Security built in (and worth explaining in interviews):
  - The password is never stored, only a salted hash of it (see set_password.py)
  - 5 wrong attempts locks sign-in for 5 minutes, like an account lockout policy
  - Every form carries a secret token, so other websites can't submit it (CSRF protection)
  - You're signed out automatically after 30 minutes without activity
  - Signed-in pages are never cached, so "Back" after signing out shows nothing
  - The report form is rate limited, so nobody can flood the queue or use up the AI quota
  - Ticket text is untrusted: it's always shown as plain text and stored using safe SQL
"""

import csv
import hmac
import io
import os
import re
import secrets
import socket
import sys
import time
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

from flask import (Flask, abort, flash, jsonify, redirect, render_template, request,
                   send_file, send_from_directory, session, url_for)
from werkzeug.security import check_password_hash

import tickets_db
import triage
# Reuse the triage script's .env reader, priority rules and Excel column names
from triage import EXCEL_COLUMNS, PRIORITIES, load_env_file

BASE_DIR = Path(__file__).resolve().parent   # the project folder


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

COMPANY_NAME = "Brightmoor IT"   # a made-up company: change it here and every page updates
PORTAL_NAME = "Desk Portal"

MAX_ATTEMPTS = 5            # wrong passwords allowed before a lockout
LOCKOUT_SECONDS = 5 * 60    # how long the lockout lasts
SESSION_MINUTES = 30        # sign out after this long without activity

REPORTS_ALLOWED = 5         # tickets one computer can send...
REPORT_WINDOW = 10 * 60     # ...in this many seconds (10 minutes)
REPORTS_PER_HOUR = 30       # most tickets the whole site accepts per hour (protects the AI quota)

# Shown at the bottom of every page, so visitors know it's a demo and nobody enters real data
AUTHOR_NAME = "Connor Williams"
PROJECT_URL = "https://github.com/Cw2854/helpdesk-ai-triage"

load_env_file(BASE_DIR / ".env")

# Set PUBLIC_SITE=1 in .env on the web host (PythonAnywhere). It turns on the extra
# protections a public website needs: HTTPS-only cookies, the HSTS header, and reading
# visitors' real IP addresses from the host's load balancer.
PUBLIC_SITE = os.environ.get("PUBLIC_SITE", "").strip() == "1"
ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "").strip()
ADMIN_PASSWORD_HASH = os.environ.get("ADMIN_PASSWORD_HASH", "").strip()
SECRET_KEY = os.environ.get("SECRET_KEY", "").strip()

if not (ADMIN_USERNAME and ADMIN_PASSWORD_HASH and SECRET_KEY):
    sys.exit("No admin login set up yet. Run:  python set_password.py")

app = Flask(__name__)
app.config.update(
    SECRET_KEY=SECRET_KEY,                  # signs the session cookie so it can't be faked
    SESSION_COOKIE_HTTPONLY=True,           # page scripts can't read the cookie
    SESSION_COOKIE_SAMESITE="Lax",          # cookie isn't sent with other sites' form posts
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),   # not used: the sign-in ends when the browser closes,
                                                     # and the 30-minute inactivity check is below
    SESSION_COOKIE_SECURE=PUBLIC_SITE,      # online, the cookie is only ever sent over HTTPS
)


# ---------------------------------------------------------------------------
# Lockout after too many wrong passwords
# Kept in memory, so restarting the app clears it. Fine for a local tool.
# ---------------------------------------------------------------------------

failed_attempts = {}   # ip address -> {"count": int, "locked_until": float}


def seconds_locked(ip):
    """How many seconds this address is still locked out for (0 = not locked)."""
    record = failed_attempts.get(ip)
    if not record:
        return 0
    remaining = record["locked_until"] - time.time()
    if remaining > 0:
        return int(remaining) + 1
    if record["locked_until"]:          # lockout has ended: start counting again
        failed_attempts.pop(ip, None)
    return 0


def record_failure(ip):
    """Count a wrong attempt. Returns how many attempts are left before a lockout."""
    record = failed_attempts.setdefault(ip, {"count": 0, "locked_until": 0})
    record["count"] += 1
    attempts_left = MAX_ATTEMPTS - record["count"]
    if attempts_left <= 0:
        record["locked_until"] = time.time() + LOCKOUT_SECONDS
    return attempts_left


# ---------------------------------------------------------------------------
# CSRF protection: a random token in the session that every form must send back
# ---------------------------------------------------------------------------

def csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return session["csrf_token"]


def check_csrf():
    sent = request.form.get("csrf_token", "")
    expected = session.get("csrf_token", "")
    if not expected or not hmac.compare_digest(sent, expected):
        abort(400, "This form has expired. Go back, refresh the page and try again.")


app.jinja_env.globals["csrf_token"] = csrf_token
app.jinja_env.globals["company_name"] = COMPANY_NAME   # available in every page template
app.jinja_env.globals["portal_name"] = PORTAL_NAME
app.jinja_env.globals["author_name"] = AUTHOR_NAME
app.jinja_env.globals["project_url"] = PROJECT_URL


def client_ip():
    """The visitor's IP address, used for the login lockout and the report form's rate limit.
    Online, every request arrives through PythonAnywhere's load balancer, so request.remote_addr
    is the balancer's address (the same for everyone). The real one is in the X-Real-IP header,
    which the balancer sets itself, so visitors can't fake it."""
    if PUBLIC_SITE:
        return request.headers.get("X-Real-IP") or request.remote_addr or "unknown"
    return request.remote_addr or "unknown"


# ---------------------------------------------------------------------------
# Sign-in checks
# ---------------------------------------------------------------------------

def signed_in(count_as_activity=True):
    """True if the admin is signed in and hasn't been inactive for too long."""
    if "user" not in session:
        return False
    if time.time() - session.get("last_active", 0) > SESSION_MINUTES * 60:
        session.clear()
        flash(f"You were signed out after {SESSION_MINUTES} minutes without activity.")
        return False
    if count_as_activity:
        session["last_active"] = time.time()
    return True


def login_required(view):
    """Put @login_required above any page that only the admin may see."""
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not signed_in():
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped


def credentials_ok(username, password):
    # The username ignores capitals ("Admin", "admin" and "ADMIN" are the same), like Windows
    # and Active Directory logins. The password is still case-sensitive.
    # compare_digest takes the same time whether the first letter or the last is wrong,
    # so attackers can't guess the username by timing the response.
    name_ok = hmac.compare_digest(username.casefold().encode(), ADMIN_USERNAME.casefold().encode())
    password_ok = check_password_hash(ADMIN_PASSWORD_HASH, password)   # always run both checks
    return name_ok and password_ok


@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; "
        "style-src 'self' https://fonts.googleapis.com; "
        "font-src https://fonts.gstatic.com; "
        "frame-ancestors 'none'"
    )
    if response.mimetype == "text/html":
        response.headers["Cache-Control"] = "no-store"
    if PUBLIC_SITE:
        # Tell browsers to only ever use HTTPS for this site from now on
        response.headers["Strict-Transport-Security"] = "max-age=31536000"
    return response


# ---------------------------------------------------------------------------
# Admin pages
# ---------------------------------------------------------------------------

@app.route("/favicon.ico")
def favicon():
    """The logo in the browser tab (browsers ask for /favicon.ico by default)."""
    return send_from_directory(app.static_folder, "favicon.ico", mimetype="image/vnd.microsoft.icon")


@app.route("/")
def home():
    """The main address shows the employee form. Admins go to /admin."""
    return redirect(url_for("report"))


# Old addresses from earlier versions, so existing bookmarks still work
@app.route("/login")
def old_login():
    return redirect(url_for("login"))


@app.route("/tickets")
def old_tickets():
    return redirect(url_for("tickets", **request.args))


@app.route("/report")
def old_report():
    return redirect(url_for("report"))


@app.route("/admin", methods=["GET", "POST"])
def login():
    # Opening the admin page always asks for the username and password, even if this browser
    # was already signed in. So if the admin leaves a computer signed in, someone else can't
    # get into the ticket queue just by clicking the admin link on the request form.
    if request.method == "GET" and signed_in(count_as_activity=False):
        session.clear()
        flash("For security, you've been signed out. Sign in again to open the ticket queue.")

    ip = client_ip()
    locked_for = seconds_locked(ip)
    username = ""
    error = None
    status = 200

    if request.method == "POST":
        check_csrf()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if locked_for:
            status = 429
        elif credentials_ok(username, password):
            failed_attempts.pop(ip, None)
            session.clear()                 # fresh session on sign-in
            session["user"] = username
            session["last_active"] = time.time()
            session.permanent = False       # the sign-in ends when the browser is closed
            return redirect(url_for("tickets"))
        else:
            attempts_left = record_failure(ip)
            if attempts_left <= 0:
                locked_for = seconds_locked(ip)
                status = 429
            else:
                plural = "attempt" if attempts_left == 1 else "attempts"
                error = (f"That username and password don't match. "
                         f"{attempts_left} {plural} left before sign-in is locked for 5 minutes.")
                status = 401

    return render_template("login.html", username=username, error=error,
                           locked_for=locked_for), status


@app.route("/admin/logout", methods=["POST"])
def logout():
    check_csrf()
    session.clear()
    flash("You've signed out.")
    return redirect(url_for("login"))


def chosen_view(value):
    """Which list to show: "open" (the default), "completed" or "deleted"."""
    return value if value in tickets_db.VIEWS else "open"


@app.route("/admin/tickets")
@login_required
def tickets():
    view = chosen_view(request.args.get("view"))

    # The priority rules from triage.py, split into ("P1", "Critical", "several users ...")
    priority_key = [(code, *text.split(" - ", 1)) for code, text in PRIORITIES.items()]

    ticket_list = [to_display(row) for row in tickets_db.tickets_in(view)]
    if view == "open":          # most urgent first
        ticket_list.sort(key=lambda t: (PRIORITY_ORDER.get(t["priority"], 9), t["id"]))
    else:                       # most recently completed / deleted first
        ticket_list.sort(key=lambda t: t["completed_at" if view == "completed" else "deleted_at"],
                         reverse=True)

    counts = {code: sum(t["priority"] == code for t in ticket_list) for code in PRIORITIES}
    return render_template(
        "tickets.html", user=session["user"], session_minutes=SESSION_MINUTES,
        priority_key=priority_key, tickets=ticket_list, counts=counts,
        review_count=sum(t["needs_review"] for t in ticket_list),
        view=view, view_counts=tickets_db.view_counts(), latest_row=tickets_db.latest_row(),
        undo=session.pop("undo", None),     # shows the "Undo" message once
    )


# ---------------------------------------------------------------------------
# Ticket actions: complete, reopen, delete, restore
# All are POST requests with the CSRF token, so a link or another website can't trigger them.
# ---------------------------------------------------------------------------

TICKET_ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]{1,40}$")   # e.g. INC-1019


def run_action(ticket_id, action):
    """Check the request, run the action, then go back to the list the admin was on."""
    check_csrf()
    if not TICKET_ID_PATTERN.match(ticket_id):
        abort(404)
    changed = action(ticket_id)
    return changed, redirect(url_for("tickets", view=chosen_view(request.form.get("view"))))


@app.route("/admin/tickets/<ticket_id>/complete", methods=["POST"])
@login_required
def complete_route(ticket_id):
    changed, back = run_action(ticket_id, lambda i: tickets_db.complete_ticket(i, session["user"]))
    if changed:
        session["undo"] = {"id": ticket_id, "action": "completed"}
    return back


@app.route("/admin/tickets/<ticket_id>/reopen", methods=["POST"])
@login_required
def reopen_route(ticket_id):
    changed, back = run_action(ticket_id, tickets_db.reopen_ticket)
    if changed:
        flash(f"{ticket_id} reopened. It's back in the Open list.")
    return back


@app.route("/admin/tickets/<ticket_id>/delete", methods=["POST"])
@login_required
def delete_route(ticket_id):
    changed, back = run_action(ticket_id, lambda i: tickets_db.delete_ticket(i, session["user"]))
    if changed:
        session["undo"] = {"id": ticket_id, "action": "deleted"}
    return back


@app.route("/admin/tickets/<ticket_id>/restore", methods=["POST"])
@login_required
def restore_route(ticket_id):
    changed, back = run_action(ticket_id, tickets_db.restore_ticket)
    if changed:
        flash(f"{ticket_id} restored.")
    return back


# ---------------------------------------------------------------------------
# Permanently deleting tickets: only from the Deleted list, and only after a confirmation page.
# This can't be undone, so it's two steps away from an open ticket.
# ---------------------------------------------------------------------------

@app.route("/admin/tickets/<ticket_id>/purge", methods=["GET", "POST"])
@login_required
def purge_route(ticket_id):
    if not TICKET_ID_PATTERN.match(ticket_id):
        abort(404)
    ticket = tickets_db.get_ticket(ticket_id)
    if not ticket or not ticket.get("deleted_at"):
        flash(f"{ticket_id} isn't in the Deleted list. Delete it first, then you can delete it permanently.")
        return redirect(url_for("tickets", view="deleted"))

    if request.method == "POST":
        check_csrf()
        if tickets_db.purge_ticket(ticket_id, by=session["user"]):
            flash(f"{ticket_id} was permanently deleted.")
        return redirect(url_for("tickets", view="deleted"))

    # GET: show the "are you sure?" page first
    return render_template("confirm_purge.html", user=session["user"],
                           tickets=[to_display(ticket)], everything=False)


@app.route("/admin/tickets/deleted/empty", methods=["GET", "POST"])
@login_required
def empty_deleted_route():
    if request.method == "POST":
        check_csrf()
        # Only the tickets listed on the confirmation page
        ids = [i for i in request.form.getlist("ticket_id") if TICKET_ID_PATTERN.match(i)]
        count = tickets_db.purge_deleted(ids, by=session["user"])
        flash(f"{count} {'ticket was' if count == 1 else 'tickets were'} permanently deleted.")
        return redirect(url_for("tickets", view="deleted"))

    deleted = [to_display(row) for row in tickets_db.tickets_in("deleted")]
    if not deleted:
        return redirect(url_for("tickets", view="deleted"))
    return render_template("confirm_purge.html", user=session["user"],
                           tickets=deleted, everything=True)


@app.route("/admin/tickets/export")
@login_required
def export_tickets():
    """Download the tickets in the current list (Open, Completed or Deleted) as Excel."""
    view = chosen_view(request.args.get("view"))
    # Leave out empty values, so tickets without expected answers aren't scored as wrong
    rows = [{k: v for k, v in row.items() if v is not None} for row in tickets_db.tickets_in(view)]
    if not rows:
        return redirect(url_for("tickets", view=view))
    buffer = io.BytesIO()
    triage.save_excel(rows, triage.build_summary(rows), buffer)   # same layout as triage.py
    buffer.seek(0)
    response = send_file(
        buffer, as_attachment=True,
        download_name=f"{view}_tickets_{datetime.now():%Y-%m-%d_%H%M}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response.headers["Cache-Control"] = "no-store"
    return response


@app.route("/admin/api/new-tickets")
def new_tickets():
    """How many tickets have arrived since the queue page was loaded. The page checks this."""
    # Checking for new tickets doesn't count as activity, so the 30-minute sign-out still works
    if not signed_in(count_as_activity=False):
        return jsonify(error="Not signed in"), 401
    since = request.args.get("since", "0")
    return jsonify(count=tickets_db.count_newer_than(int(since) if since.isdigit() else 0))


# ---------------------------------------------------------------------------
# Employee page: report a problem
# No sign-in needed: in a real company this would sit on the internal network.
# ---------------------------------------------------------------------------

recent_reports = {}   # ip address -> times it sent a ticket (for rate limiting)
all_reports = []      # times any ticket was sent (for the whole-site hourly cap)
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def too_many_reports(ip):
    now = time.time()
    times = [t for t in recent_reports.get(ip, []) if now - t < REPORT_WINDOW]
    recent_reports[ip] = times
    return len(times) >= REPORTS_ALLOWED


def site_too_busy():
    """True if the whole site has had REPORTS_PER_HOUR tickets in the last hour."""
    now = time.time()
    all_reports[:] = [t for t in all_reports if now - t < 3600]
    return len(all_reports) >= REPORTS_PER_HOUR


def check_report(form):
    """Check the form. Returns a dictionary of {field: problem}; empty means it's all fine."""
    problems = {}
    if not form["name"]:
        problems["name"] = "Enter your name."
    elif len(form["name"]) > 80:
        problems["name"] = "Use 80 characters or fewer."
    if not EMAIL_PATTERN.match(form["email"]) or len(form["email"]) > 120:
        problems["email"] = "Enter your work email, like name@company.com."
    if len(form["subject"]) < 3:
        problems["subject"] = "Give the problem a short title, like 'Laptop won't turn on'."
    elif len(form["subject"]) > 120:
        problems["subject"] = "Keep the title to 120 characters or fewer."
    if len(form["details"]) < 10:
        problems["details"] = "Describe the problem in a sentence or two."
    elif len(form["details"]) > 4000:
        problems["details"] = "Keep the details to 4,000 characters or fewer."
    return problems


def ai_triage(ticket):
    """Ask the AI to triage a new ticket. Returns the triage fields to save with it."""
    try:
        settings = triage.get_settings()
    except SystemExit as error:          # triage.py stops when the API key is missing
        raise RuntimeError(str(error))
    result = triage.triage_ticket(settings, {
        "subject": ticket["subject"], "from_name": ticket["from_name"], "body": ticket["message"],
    })
    result["model"] = settings["model"]
    return result


@app.route("/requests", methods=["GET", "POST"])
def report():
    form = {"name": "", "email": "", "subject": "", "details": ""}
    problems, notice, status = {}, None, 200

    if request.method == "POST":
        check_csrf()
        # Spam trap: a hidden box that people never see but spam bots fill in.
        # If it's filled in, quietly do nothing.
        if request.form.get("website"):
            return redirect(url_for("report"))

        form = {key: request.form.get(key, "").strip() for key in form}
        problems = check_report(form)
        ip = client_ip()

        if not problems and too_many_reports(ip):
            notice = "You've sent several tickets in the last few minutes. Wait a little, then try again."
            status = 429
        elif not problems and site_too_busy():
            notice = "The service desk is getting a lot of tickets right now. Try again in a few minutes."
            status = 429
        elif not problems:
            recent_reports.setdefault(ip, []).append(time.time())
            all_reports.append(time.time())
            ticket = {
                "received": datetime.now().strftime("%Y-%m-%d %H:%M"),
                "from_name": form["name"], "from_email": form["email"],
                "subject": form["subject"], "message": form["details"],
                "source": "Report form",
            }
            # Let the AI triage it. If that fails, save it anyway and flag it for a person:
            # a ticket must never be lost just because the AI is unavailable.
            try:
                ticket.update(ai_triage(ticket))
            except Exception as error:
                ticket.update({
                    "category": "Error", "priority": "", "needs_human_review": "Yes",
                    "review_reason": f"AI triage failed, so triage this one by hand ({error})",
                })
            ticket["triaged_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
            session["sent_ticket"] = {"id": tickets_db.create_ticket(ticket), "email": form["email"]}
            return redirect(url_for("report_sent"))      # so refreshing can't send it twice
        else:
            status = 400

    return render_template("report.html", form=form, problems=problems, notice=notice), status


@app.route("/requests/sent")
def report_sent():
    sent = session.pop("sent_ticket", None)
    if not sent:
        return redirect(url_for("report"))
    return render_template("report_sent.html", ticket_id=sent["id"], email=sent["email"])


# ---------------------------------------------------------------------------
# Turning stored tickets into what the queue page shows
# ---------------------------------------------------------------------------

PRIORITY_ORDER = {"P1": 1, "P2": 2, "P3": 3, "P4": 4}


def tidy_steps(text):
    """Put each numbered step on its own line, if the AI wrote them all on one line."""
    if "\n" not in text and text.startswith("1."):
        text = re.sub(r"\s+(?=[2-9]\.\s+[A-Z])", "\n", text)
    return text


def to_display(row):
    def text(field):
        value = row.get(field)
        return "" if value is None else str(value).strip()

    try:
        confidence = f"{float(row.get('confidence')):.0%}"
    except (TypeError, ValueError):
        confidence = ""

    received = text("received")
    try:
        received_label = datetime.strptime(received, "%Y-%m-%d %H:%M").strftime("%d %b, %H:%M")
    except ValueError:
        received_label = received

    priority = text("priority").upper()
    return {
        "id": text("ticket_id"),
        "priority": priority if priority in PRIORITY_ORDER else "",
        "category": text("category"),
        "issue": text("issue_summary") or text("subject"),
        "subject": text("subject"),
        "message": text("message"),
        "from_name": text("from_name"),
        "from_email": text("from_email"),
        "received": received,
        "received_label": received_label,
        "suggested_fix": tidy_steps(text("suggested_fix")),
        "draft_reply": text("draft_reply"),
        "needs_review": text("needs_human_review") == "Yes",
        "review_reason": text("review_reason"),
        "confidence": confidence,
        "ai_reason": text("ai_reason"),
        "expected_category": text("expected_category"),
        "category_correct": text("category_correct"),
        "expected_priority": text("expected_priority"),
        "priority_correct": text("priority_correct"),
        "source": text("source"),
        "completed_at": text("completed_at"),
        "completed_label": nice_time(text("completed_at")),
        "completed_by": text("completed_by"),
        "deleted_at": text("deleted_at"),
        "deleted_label": nice_time(text("deleted_at")),
        "deleted_by": text("deleted_by"),
    }


def nice_time(value):
    """Turn "2026-10-24 16:12:34" into "24 Oct, 16:12"."""
    try:
        return datetime.strptime(value[:16], "%Y-%m-%d %H:%M").strftime("%d %b, %H:%M")
    except ValueError:
        return value


# ---------------------------------------------------------------------------
# Bringing in results from triage.py
# ---------------------------------------------------------------------------

def latest_results_file():
    """The most recent triaged_tickets file in the project folder (Excel or CSV), or None."""
    files = [f for pattern in ("triaged_tickets*.xlsx", "triaged_tickets*.csv")
             for f in BASE_DIR.glob(pattern)
             if not f.name.startswith("~$")]          # skip Excel's temporary lock files
    return max(files, key=lambda f: f.stat().st_mtime, default=None)


def read_rows(path):
    """Read every row of the results file as a dictionary of {field name: value}."""
    # The Excel file uses friendly headings ("Suggested fix (for the analyst)"),
    # so turn them back into the field names used in triage.py ("suggested_fix").
    heading_to_field = {heading: field for field, heading, _, _ in EXCEL_COLUMNS}

    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))

    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Tickets"] if "Tickets" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headings = [heading_to_field.get(h, h) for h in next(rows, [])]
        return [dict(zip(headings, row)) for row in rows if any(row)]
    finally:
        workbook.close()


def import_latest_results():
    """Copy tickets from the newest triaged_tickets file into the database.
    Tickets already in the database (same ticket number) are skipped, so this is safe to repeat."""
    path = latest_results_file()
    if not path:
        return 0
    added = 0
    try:
        for row in read_rows(path):
            for key, value in row.items():
                if isinstance(value, datetime):          # Excel sometimes turns text into dates
                    row[key] = value.strftime("%Y-%m-%d %H:%M")
            row["ticket_id"] = str(row.get("ticket_id") or "").strip()
            if row["ticket_id"]:
                row["source"] = f"triage.py results ({path.name})"
                added += tickets_db.add_ticket(row)
    except PermissionError:
        print(f"Couldn't read {path.name} because it's open in another program. Close it and restart.")
    return added


@app.errorhandler(400)
def bad_request(error):
    return render_template("error.html", message=error.description), 400


tickets_db.create_table()
IMPORTED_AT_START = import_latest_results()     # copies in tickets from the latest Excel results


def already_running(port):
    """True if something is already answering on this port, such as an old copy of the app.
    (Windows lets two programs listen on the same port, and the old one keeps answering.)"""
    try:
        with socket.create_connection(("127.0.0.1", port), timeout=0.5):
            return True
    except OSError:
        return False


if __name__ == "__main__":
    if already_running(5000):
        sys.exit(
            "\nAn old copy of the app is still running on port 5000, so your browser would keep\n"
            "showing the old version. Stop it first by running this command, then start the app again:\n\n"
            "    taskkill /F /IM python.exe\n")
    if IMPORTED_AT_START:
        print(f"\nAdded {IMPORTED_AT_START} tickets from your latest triage results to the database.")
    print("\nDesk Portal running.")
    print("  Admin login:  http://127.0.0.1:5000/admin")
    print("  Requests:     http://127.0.0.1:5000/requests")
    print("Press Ctrl+C here to stop it.\n")
    # 127.0.0.1 = only this computer can reach it. debug stays off so errors never show code.
    app.run(host="127.0.0.1", port=5000, debug=False)
